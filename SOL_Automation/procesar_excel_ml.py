"""
Script para procesar Excel ML_pendientes_test y actualizar documentos Word
Integra la automatización de SOL para generar evidencias
"""
import openpyxl
from openpyxl import load_workbook
from docx import Document
from docx.shared import Inches
import os
import sys
from pathlib import Path

# Agregar la ruta del proyecto principal para importar módulos
sys.path.append(r'C:\Users\eliseo_lopezp\proyecto1\automatizacion_web')

# Importar módulos necesarios de la automatización
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from SOL.SOL_SKU import SOLSKUClass

def iniciar_sesion_sol():
    """
    Inicia sesión en SOL y retorna el driver
    
    Returns:
        driver: WebDriver de Chrome con sesión iniciada
    """
    try:
        print("  🤖 Iniciando navegador y sesión...")
        
        # Configurar Chrome
        chrome_options = Options()
        chrome_options.add_argument('--start-maximized')
        
        # Iniciar el navegador
        driver = webdriver.Chrome(options=chrome_options)
        
        # Importar y ejecutar el login
        from Page.Page_Login import LoginPage
        from Page.page_menu import MenuPage
        
        # Navegar a la URL
        driver.get("http://was7tr1.siman.com/AccesoSUMMER/")
        
        # Login
        login_page = LoginPage(driver)
        print("  📝 Iniciando sesión...")
        login_page.login("ECHAVEZ", "Noviembre2025*")
        
        import time
        time.sleep(3)
        
        # Acceder al menú SOL
        menu_page = MenuPage(driver)
        print("  📂 Accediendo a SOL...")
        menu_page.acceder_frame_menu()
        menu_page.hacer_click_en_summer()
        menu_page.click_SOL()
        
        time.sleep(3)
        
        print("  ✅ Sesión iniciada correctamente")
        return driver
        
    except Exception as e:
        print(f"  ❌ Error al iniciar sesión: {e}")
        if driver:
            driver.quit()
        return None

def procesar_sku_en_sesion(driver, sku_numero, carpeta_evidencias):
    """
    Procesa un SKU usando una sesión ya iniciada
    
    Args:
        driver: WebDriver con sesión activa
        sku_numero: Número de SKU a procesar
        carpeta_evidencias: Carpeta donde guardar las capturas
    
    Returns:
        bool: True si la prueba fue válida, False en caso contrario
    """
    try:
        print(f"  🔍 Procesando SKU {sku_numero}...")
        
        # Acceder a Mantenimiento de SKU
        from SOL.SOL_SKU import ejecutar_busqueda_sku_sol
        
        if not ejecutar_busqueda_sku_sol(driver):
            print(f"  ❌ No se pudo acceder a Mantenimiento de Item")
            return False
        
        # Buscar el SKU específico usando procesar_lista_skus
        print(f"  🔍 Buscando item {sku_numero}...")
        sol_sku = SOLSKUClass(driver)
        sol_sku.carpeta_actual = carpeta_evidencias
        
        # Desactivar generación de documento Word (lo haremos manualmente después)
        generar_doc_original = sol_sku.generar_documento_evidencias
        sol_sku.generar_documento_evidencias = lambda x: print("  📝 Omitiendo generación de documento Word automático...")
        
        resultados = sol_sku.procesar_lista_skus([sku_numero], intervalo_segundos=2)
        
        # Restaurar método original
        sol_sku.generar_documento_evidencias = generar_doc_original
        
        # Verificar si la prueba fue válida (si hay datos en la tabla Donde)
        prueba_valida = sol_sku.prueba_valida if hasattr(sol_sku, 'prueba_valida') else True
        
        print(f"  ✅ SKU {sku_numero} procesado correctamente")
        return prueba_valida
        
    except Exception as e:
        print(f"  ❌ Error al procesar SKU {sku_numero}: {e}")
        return False

def obtener_skus_del_excel(ruta_excel):
    """
    Lee el Excel y obtiene todos los SKUs diferentes de 0
    
    Returns:
        list: Lista de diccionarios con información de cada SKU
              {'sku': '123456', 'hipervínculo': 'ruta/documento.docx', 'fila': 2}
    """
    try:
        wb = load_workbook(ruta_excel)
        ws = wb.active
        
        skus_encontrados = []
        
        # Buscar la columna SKU (asumiendo que está en alguna de las primeras columnas)
        headers = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(1, col).value
            if cell_value:
                headers.append(str(cell_value).strip().upper())
            else:
                headers.append("")
        
        print(f"📋 Columnas encontradas en el Excel:")
        for idx, header in enumerate(headers, 1):
            if header:
                print(f"   Columna {idx}: {header}")
        print()
        
        # Encontrar índice de columna SKU
        sku_col = None
        hyperlink_col = None
        
        for idx, header in enumerate(headers, 1):
            if 'SKU' in header:
                sku_col = idx
                print(f"✅ Columna SKU encontrada en posición {idx}: {ws.cell(1, idx).value}")
            if 'RUTA' in header and 'EVIDENCIA' in header:
                hyperlink_col = idx
                print(f"✅ Columna Ruta Evidencia encontrada en posición {idx}: {ws.cell(1, idx).value}")
        
        if not sku_col:
            print("❌ No se encontró columna SKU")
            print("   Verifica que el Excel tenga una columna llamada 'SKU'")
            wb.close()
            return []
        
        if not hyperlink_col:
            print("❌ No se encontró columna Ruta Evidencia")
            print("   Verifica que el Excel tenga una columna llamada 'Ruta Evidencia'")
            wb.close()
            return []
        
        # Recorrer filas buscando SKUs diferentes de 0
        for row in range(2, ws.max_row + 1):
            sku_cell = ws.cell(row, sku_col)
            sku_value = sku_cell.value
            
            # Verificar si el SKU es diferente de 0 y no está vacío
            if sku_value and str(sku_value).strip() not in ['0', '', 'None']:
                sku_info = {
                    'sku': str(sku_value).strip(),
                    'fila': row,
                    'hipervínculo': None
                }
                
                # Buscar hipervínculo en la misma fila
                if hyperlink_col:
                    hyperlink_cell = ws.cell(row, hyperlink_col)
                    hyperlink_value = None
                    
                    # Intentar obtener hipervínculo de la celda
                    if hyperlink_cell.hyperlink:
                        hyperlink_value = hyperlink_cell.hyperlink.target
                    elif hyperlink_cell.value:
                        # Si es una fórmula HYPERLINK, extraer la ruta
                        valor_celda = str(hyperlink_cell.value)
                        if 'HYPERLINK' in valor_celda:
                            # Extraer la ruta entre comillas: =HYPERLINK("ruta", "texto")
                            import re
                            match = re.search(r'HYPERLINK\("([^"]+)"', valor_celda)
                            if match:
                                hyperlink_value = match.group(1)
                        else:
                            hyperlink_value = valor_celda
                    
                    sku_info['hipervínculo'] = hyperlink_value
                
                skus_encontrados.append(sku_info)
                print(f"  📌 Fila {row}: SKU={sku_info['sku']}, Hipervínculo={sku_info['hipervínculo']}")
        
        wb.close()
        print(f"\n✅ Total de SKUs encontrados: {len(skus_encontrados)}")
        return skus_encontrados
        
    except Exception as e:
        print(f"❌ Error al leer Excel: {e}")
        return []

def actualizar_fila_excel(ruta_excel, fila, nombre_estado="PASS", comentario="Prueba exitosa", flag="Y"):
    """
    Actualiza las columnas Nombre Estado, Comentario y Flag_process en el Excel
    
    Args:
        ruta_excel: Ruta al archivo Excel
        fila: Número de fila a actualizar
        nombre_estado: Valor para columna Nombre Estado (default: "PASS")
        comentario: Valor para columna Comentario (default: "Prueba exitosa")
        flag: Valor para columna Flag_process (default: "Y")
    """
    try:
        wb = load_workbook(ruta_excel)
        ws = wb.active
        
        # Buscar columnas Nombre Estado, Comentario y Flag_process
        headers = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(1, col).value
            if cell_value:
                headers.append((col, str(cell_value).strip().upper()))
            else:
                headers.append((col, ""))
        
        nombre_estado_col = None
        comentario_col = None
        flag_col = None
        
        for col_idx, header in headers:
            if 'NOMBRE' in header and 'ESTADO' in header:
                nombre_estado_col = col_idx
            if 'COMENTARIO' in header:
                comentario_col = col_idx
            if 'FLAG' in header and 'PROCESS' in header:
                flag_col = col_idx
        
        # Actualizar celdas
        if nombre_estado_col:
            ws.cell(fila, nombre_estado_col).value = nombre_estado
            print(f"    ✅ Nombre Estado actualizado a: {nombre_estado}")
        else:
            print(f"    ⚠️ No se encontró columna 'Nombre Estado'")
        
        if comentario_col:
            ws.cell(fila, comentario_col).value = comentario
            print(f"    ✅ Comentario actualizado a: {comentario}")
        else:
            print(f"    ⚠️ No se encontró columna 'Comentario'")
        
        if flag_col:
            ws.cell(fila, flag_col).value = flag
            print(f"    ✅ Flag_process actualizado a: {flag}")
        else:
            print(f"    ⚠️ No se encontró columna 'Flag_process'")
        
        wb.save(ruta_excel)
        wb.close()
        return True
        
    except Exception as e:
        print(f"    ❌ Error al actualizar Excel: {e}")
        return False

def actualizar_documento_word(ruta_word, sku_numero, carpeta_evidencias, prueba_valida=True):
    """
    Actualiza el documento Word:
    - Elimina contenido desde "Pasos de Prueba" hasta "Resultado Esperado"
    - Inserta las 6 capturas del proceso
    
    Args:
        ruta_word: Ruta al documento Word
        sku_numero: Número del SKU
        carpeta_evidencias: Carpeta donde están las capturas
        prueba_valida: Si la prueba fue exitosa (hay datos en tabla Donde)
    """
    try:
        if not os.path.exists(ruta_word):
            print(f"  ❌ Archivo Word no encontrado: {ruta_word}")
            return False
        
        doc = Document(ruta_word)
        
        # Buscar la sección "Pasos de Prueba" y "Resultado Esperado"
        indice_pasos_prueba = None
        indice_resultado_esperado = None
        contenido_despues = []  # Guardar todo después de "Pasos de Prueba"
        
        for idx, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            
            # Encontrar "Pasos de Prueba:" o "Pasos de la Prueba:"
            if not indice_pasos_prueba and ('Pasos de Prueba' in text or 'Pasos de la Prueba' in text):
                indice_pasos_prueba = idx
                print(f"  🔍 Encontrada sección 'Pasos de Prueba' en párrafo {idx}: '{text}'")
            
            # Encontrar "Resultado Esperado:" y guardar todo desde ahí
            elif indice_pasos_prueba and 'Resultado Esperado' in text:
                indice_resultado_esperado = idx
                print(f"  🔍 Encontrada sección 'Resultado Esperado' en párrafo {idx}: '{text}'")
                # Guardar "Resultado Esperado" y todo lo que sigue
                for i in range(idx, len(doc.paragraphs)):
                    contenido_despues.append(doc.paragraphs[i])
                break
        
        if indice_pasos_prueba is not None and indice_resultado_esperado is not None:
            # Eliminar todo después de "Pasos de Prueba"
            num_a_eliminar = len(doc.paragraphs) - indice_pasos_prueba - 1
            print(f"  🗑️ Eliminando {num_a_eliminar} párrafos después de 'Pasos de Prueba'")
            # Eliminar en orden inverso
            for idx in range(len(doc.paragraphs) - 1, indice_pasos_prueba, -1):
                p = doc.paragraphs[idx]
                p._element.getparent().remove(p._element)
            
            print(f"  📝 Insertando evidencias después del párrafo {indice_pasos_prueba}")
        else:
            print(f"  ⚠️ No se encontraron ambas secciones correctamente")
            if not indice_pasos_prueba:
                print(f"  ℹ️ No se encontró 'Pasos de Prueba', agregando al final")
                indice_pasos_prueba = len(doc.paragraphs) - 1
        
        # Buscar las 6 capturas en la carpeta de evidencias
        capturas = [
            f"1_mantenimiento_inicial_{sku_numero}.png",
            f"2_sku_ingresado_{sku_numero}.png",
            f"3_tabla_resultados_{sku_numero}.png",
            f"4_opciones_location_desplegadas_{sku_numero}.png",
            f"5_segunda_pestana_donde_{sku_numero}.png",
            f"6_validacion_compania_{sku_numero}.png"
        ]
        
        # Insertar las evidencias después de "Pasos de Prueba"
        # Usamos add_paragraph que agrega al final, ya que limpiamos el contenido anterior
        
        descripciones_pasos = [
            "Paso 1: Ingreso a SOL - Mantenimiento de Item",
            "Paso 2: Item ingresado en el campo de búsqueda",
            "Paso 3: Selección de una compañía",
            "Paso 4: Click en pestaña Donde (primera vez)",
            "Paso 5: Click en pestaña Donde (segunda vez)",
            "Paso 6: Validación - Compañía mostrada en la tabla"
        ]
        
        # Insertar las evidencias después de "Pasos de Prueba"
        for i, nombre_captura in enumerate(capturas, 1):
            ruta_captura = os.path.join(carpeta_evidencias, nombre_captura)
            
            if os.path.exists(ruta_captura):
                # Agregar descripción del paso
                p = doc.add_paragraph()
                run = p.add_run(descripciones_pasos[i-1])
                run.bold = True
                
                # Agregar imagen
                doc.add_picture(ruta_captura, width=Inches(6.0))
                
                # Espacio entre pasos
                doc.add_paragraph()
                
                print(f"    ✅ Captura {i} insertada: {descripciones_pasos[i-1]}")
            else:
                print(f"    ⚠️ Captura no encontrada: {nombre_captura}")
        
        # Agregar mensaje de resultado
        doc.add_paragraph()
        resultado_p = doc.add_paragraph()
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
        resultado_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        if prueba_valida:
            resultado_run = resultado_p.add_run("✓ RESULTADO: PRUEBA EXITOSA")
            resultado_run.font.bold = True
            resultado_run.font.size = Pt(14)
            resultado_run.font.color.rgb = RGBColor(0, 128, 0)
        else:
            resultado_run = resultado_p.add_run("✗ RESULTADO: PRUEBA NO VÁLIDA - Sin datos en tabla")
            resultado_run.font.bold = True
            resultado_run.font.size = Pt(14)
            resultado_run.font.color.rgb = RGBColor(255, 0, 0)
        
        doc.add_paragraph()
        
        # Restaurar "Resultado Esperado" y todo lo que sigue
        if contenido_despues:
            print(f"  📝 Restaurando 'Resultado Esperado' y contenido posterior ({len(contenido_despues)} párrafos)")
            for para in contenido_despues:
                # Copiar el formato y contenido del párrafo
                new_para = doc.add_paragraph(para.text)
                # Copiar formato
                new_para.paragraph_format.alignment = para.paragraph_format.alignment
                # Copiar runs con formato
                new_para.clear()
                for run in para.runs:
                    new_run = new_para.add_run(run.text)
                    new_run.bold = run.bold
                    new_run.italic = run.italic
                    new_run.underline = run.underline
                    if run.font.size:
                        new_run.font.size = run.font.size
        
        # Guardar documento
        doc.save(ruta_word)
        print(f"  ✅ Documento actualizado: {ruta_word}")
        
        # Eliminar las imágenes PNG después de insertarlas en el Word
        print(f"  🗑️ Eliminando capturas PNG...")
        for nombre_captura in capturas:
            ruta_captura = os.path.join(carpeta_evidencias, nombre_captura)
            if os.path.exists(ruta_captura):
                try:
                    os.remove(ruta_captura)
                    print(f"    ✅ Eliminada: {nombre_captura}")
                except Exception as e:
                    print(f"    ⚠️ No se pudo eliminar {nombre_captura}: {e}")
        
        return True
        
    except Exception as e:
        print(f"  ❌ Error al actualizar Word: {e}")
        return False

def main():
    print("=" * 70)
    print("PROCESAMIENTO DE EXCEL ML_pendientes_test")
    print("=" * 70)
    
    # Rutas
    ruta_excel = r"C:\Users\eliseo_lopezp\scale001\Documentos\ML_pendientes_test.xlsx"
    carpeta_evidencias = r"C:\Users\eliseo_lopezp\proyecto1\automatizacion_web\Evidencias"
    
    # Verificar que existe el Excel
    if not os.path.exists(ruta_excel):
        print(f"❌ Excel no encontrado: {ruta_excel}")
        return
    
    print(f"📂 Excel: {ruta_excel}")
    print(f"📂 Carpeta Evidencias: {carpeta_evidencias}\n")
    
    # Obtener SKUs del Excel
    print("🔍 Leyendo Excel y buscando SKUs...\n")
    skus = obtener_skus_del_excel(ruta_excel)
    
    if not skus:
        print("❌ No se encontraron SKUs para procesar")
        return
    
    # Procesar cada SKU
    print(f"\n{'=' * 70}")
    print(f"PROCESANDO SKUs CON AUTOMATIZACIÓN")
    print(f"{'=' * 70}\n")
    
    # Iniciar sesión una sola vez
    driver = iniciar_sesion_sol()
    if not driver:
        print("❌ No se pudo iniciar sesión en SOL")
        return
    
    try:
        for idx, sku_info in enumerate(skus, 1):
            print(f"\n{'=' * 70}")
            print(f"--- SKU {idx}/{len(skus)} ---")
            print(f"  SKU: {sku_info['sku']}")
            print(f"  Fila: {sku_info['fila']}")
            print(f"{'=' * 70}")
            
            # Paso 1: Procesar SKU en la sesión existente
            prueba_valida = procesar_sku_en_sesion(driver, sku_info['sku'], carpeta_evidencias)
            if prueba_valida is not False:  # True o cualquier valor que no sea False
                print(f"  ✅ Evidencias generadas correctamente")
                
                # Paso 2: Actualizar documento Word
                if sku_info['hipervínculo']:
                    print(f"  Hipervínculo: {sku_info['hipervínculo']}")
                    
                    # Construir ruta completa al documento Word
                    carpeta_base = os.path.dirname(ruta_excel)
                    ruta_completa_word = os.path.join(carpeta_base, sku_info['hipervínculo'])
                    ruta_completa_word = os.path.normpath(ruta_completa_word)
                    
                    print(f"  📄 Actualizando documento Word...")
                    if actualizar_documento_word(ruta_completa_word, sku_info['sku'], carpeta_evidencias, prueba_valida):
                        # Paso 3: Actualizar Excel según resultado de la prueba
                        print(f"  📝 Actualizando Excel...")
                        if prueba_valida:
                            actualizar_fila_excel(ruta_excel, sku_info['fila'], "PASS", "Prueba exitosa", "Y")
                        else:
                            actualizar_fila_excel(ruta_excel, sku_info['fila'], "FAIL", "Prueba no exitosa - Sin datos en tabla", "Y")
                        print(f"  ✅ SKU {sku_info['sku']} procesado completamente")
                    else:
                        print(f"  ⚠️ No se pudo actualizar el documento Word")
                else:
                    print(f"  ⚠️ No hay hipervínculo para este SKU")
            else:
                print(f"  ❌ Error al generar evidencias para SKU {sku_info['sku']}")
                print(f"  ⏭️ Continuando con el siguiente SKU...")
            
            print()
    
    finally:
        # Cerrar el navegador al terminar todos los SKUs
        if driver:
            driver.quit()
            print("\n🔒 Navegador cerrado")
    
    print("=" * 70)
    print("✅ PROCESO COMPLETADO")
    print("=" * 70)

if __name__ == "__main__":
    main()
