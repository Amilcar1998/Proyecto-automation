import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


class GeneradorDocumento:
    FUENTE = "Aptos"
    COLOR_TITULO = RGBColor(31, 41, 55)
    COLOR_SUBTITULO = RGBColor(55, 65, 81)
    COLOR_RESULTADO = {
        "PASS": RGBColor(22, 101, 52),
        "FAIL": RGBColor(185, 28, 28),
        "SKIP": RGBColor(146, 64, 14),
        "ERROR": RGBColor(127, 29, 29),
    }

    def __init__(self):
        self.logger = logging.getLogger("GeneradorDocumento")
        self.doc = Document()
        self.base_actual = None

        self.logger.info("INICIO crear_documento")
        self._configurar_estilos_documento()
        self._crear_encabezado()
        self.logger.info("FIN crear_documento")

    def _crear_encabezado(self):
        titulo = self.doc.add_heading("", 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        titulo.paragraph_format.space_after = Pt(2)
        run_titulo = titulo.add_run("Evidencia de Automatizacion de Cambios de Precio")
        self._aplicar_fuente_run(run_titulo, self.FUENTE, 18, bold=True, color=self.COLOR_TITULO)

        subtitulo = self.doc.add_paragraph()
        subtitulo.paragraph_format.space_after = Pt(10)
        run_subtitulo = subtitulo.add_run(f"Fecha de ejecucion: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        self._aplicar_fuente_run(run_subtitulo, self.FUENTE, 10, color=self.COLOR_SUBTITULO)

    def _configurar_estilos_documento(self):
        self.logger.info("INICIO _configurar_estilos_documento")

        try:
            estilo_normal = self.doc.styles["Normal"]
            estilo_normal.font.name = self.FUENTE
            estilo_normal.font.size = Pt(11)
            estilo_normal._element.rPr.rFonts.set(qn("w:ascii"), self.FUENTE)
            estilo_normal._element.rPr.rFonts.set(qn("w:hAnsi"), self.FUENTE)
            estilo_normal._element.rPr.rFonts.set(qn("w:eastAsia"), self.FUENTE)
            estilo_normal._element.rPr.rFonts.set(qn("w:cs"), self.FUENTE)
            estilo_normal.paragraph_format.space_before = Pt(0)
            estilo_normal.paragraph_format.space_after = Pt(3)
            estilo_normal.paragraph_format.line_spacing = 1
        except Exception as exc:
            self.logger.warning(f"No se pudo configurar estilo Normal: {exc}")

        try:
            for nombre_estilo in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
                if nombre_estilo not in self.doc.styles:
                    continue
                estilo = self.doc.styles[nombre_estilo]
                estilo.font.name = self.FUENTE
                estilo._element.rPr.rFonts.set(qn("w:ascii"), self.FUENTE)
                estilo._element.rPr.rFonts.set(qn("w:hAnsi"), self.FUENTE)
                estilo._element.rPr.rFonts.set(qn("w:eastAsia"), self.FUENTE)
                estilo._element.rPr.rFonts.set(qn("w:cs"), self.FUENTE)
                if nombre_estilo == "Title":
                    estilo.font.size = Pt(18)
                elif nombre_estilo == "Heading 1":
                    estilo.font.size = Pt(14)
                elif nombre_estilo == "Heading 2":
                    estilo.font.size = Pt(12)
                else:
                    estilo.font.size = Pt(11)
        except Exception as exc:
            self.logger.warning(f"No se pudieron ajustar headings: {exc}")

        self.logger.info("FIN _configurar_estilos_documento")

    def _aplicar_fuente_run(self, run, nombre_fuente="Aptos", tamano=11, bold=False, color=None):
        try:
            run.font.name = nombre_fuente
            run.font.size = Pt(tamano)
            run.bold = bold
            if color is not None:
                run.font.color.rgb = color
            run._element.rPr.rFonts.set(qn("w:ascii"), nombre_fuente)
            run._element.rPr.rFonts.set(qn("w:hAnsi"), nombre_fuente)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), nombre_fuente)
            run._element.rPr.rFonts.set(qn("w:cs"), nombre_fuente)
        except Exception as exc:
            self.logger.warning(f"No se pudo aplicar fuente al run: {exc}")

    def _agregar_parrafo_con_fuente(self, texto, tamano=11, bold=False, color=None):
        parrafo = self.doc.add_paragraph()
        parrafo.paragraph_format.space_before = Pt(0)
        parrafo.paragraph_format.space_after = Pt(3)
        parrafo.paragraph_format.line_spacing = 1
        run = parrafo.add_run(str(texto))
        self._aplicar_fuente_run(run, self.FUENTE, tamano, bold=bold, color=color)
        return parrafo

    def _set_cell_shading(self, cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def _aplicar_formato_celda(self, cell, tamano=10, bold=False):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for parrafo in cell.paragraphs:
            parrafo.paragraph_format.space_before = Pt(0)
            parrafo.paragraph_format.space_after = Pt(0)
            parrafo.paragraph_format.line_spacing = 1
            for run in parrafo.runs:
                self._aplicar_fuente_run(
                    run,
                    self.FUENTE,
                    tamano,
                    bold=bold,
                    color=self.COLOR_TITULO if bold else None,
                )

    def iniciar_cp(self, nombre_cp):
        self.logger.info(f"INICIO iniciar_cp | cp={nombre_cp}")
        heading = self.doc.add_heading(level=1)
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(4)
        heading.paragraph_format.line_spacing = 1
        run = heading.add_run(str(nombre_cp))
        self._aplicar_fuente_run(run, self.FUENTE, 14, bold=True, color=self.COLOR_TITULO)
        self.logger.info("FIN iniciar_cp")

    def establecer_base(self, base):
        if self.base_actual == base:
            return
        self.base_actual = base
        self._agregar_parrafo_con_fuente(f"Base procesada: {base}", 11, bold=True, color=self.COLOR_SUBTITULO)

    def iniciar_etapa(self, nombre_etapa):
        self.logger.info(f"INICIO iniciar_etapa | etapa={nombre_etapa}")
        heading = self.doc.add_heading(level=1)
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(4)
        heading.paragraph_format.line_spacing = 1
        run = heading.add_run(str(nombre_etapa))
        self._aplicar_fuente_run(run, self.FUENTE, 14, bold=True, color=self.COLOR_TITULO)
        self.logger.info("FIN iniciar_etapa")

    def iniciar_alias(self, nombre_alias):
        self.logger.info(f"INICIO iniciar_alias | alias={nombre_alias}")
        heading = self.doc.add_heading(level=2)
        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.space_after = Pt(2)
        heading.paragraph_format.line_spacing = 1
        run = heading.add_run(str(nombre_alias))
        self._aplicar_fuente_run(run, self.FUENTE, 12, bold=True, color=self.COLOR_SUBTITULO)
        self.logger.info("FIN iniciar_alias")

    def agregar_texto(self, texto):
        self.logger.info(f"INICIO agregar_texto | texto={texto}")
        self._agregar_parrafo_con_fuente(texto, 11)
        self.logger.info("FIN agregar_texto")

    def agregar_tabla(self, datos, encabezados=None):
        self.logger.info(f"INICIO agregar_tabla | registros={len(datos) if datos else 0}")

        if not datos:
            self._agregar_parrafo_con_fuente("Sin registros encontrados", 11, color=self.COLOR_SUBTITULO)
            self.logger.info("FIN agregar_tabla | sin datos")
            return

        try:
            primera_fila = datos[0]
            columnas = len(primera_fila)
        except Exception:
            self._agregar_parrafo_con_fuente("No fue posible construir la tabla con los datos recibidos.", 11)
            self.logger.info("FIN agregar_tabla | error estructura")
            return

        tabla = self.doc.add_table(rows=1, cols=columnas)
        tabla.style = "Table Grid"
        tabla.autofit = True

        for i in range(columnas):
            celda = tabla.rows[0].cells[i]
            celda.text = encabezados[i] if encabezados and i < len(encabezados) else f"COL{i+1}"
            self._set_cell_shading(celda, "E5E7EB")
            self._aplicar_formato_celda(celda, tamano=10, bold=True)

        for fila in datos:
            row = tabla.add_row().cells
            for i in range(columnas):
                valor = ""
                try:
                    if i < len(fila):
                        valor = "" if fila[i] is None else str(fila[i])
                except Exception:
                    valor = ""
                row[i].text = valor
                self._aplicar_formato_celda(row[i], tamano=10, bold=False)

        self.logger.info("FIN agregar_tabla")

    def agregar_resultado(self, resultado):
        self.logger.info(f"INICIO agregar_resultado | resultado={resultado}")
        color = self.COLOR_RESULTADO.get(str(resultado).strip().upper(), self.COLOR_TITULO)
        self._agregar_parrafo_con_fuente(f"RESULTADO: {resultado}", 11, bold=True, color=color)
        self.logger.info("FIN agregar_resultado")

    def guardar(self, ruta_archivo):
        self.logger.info(f"INICIO guardar | ruta={ruta_archivo}")
        try:
            ruta = Path(ruta_archivo)
            ruta.parent.mkdir(parents=True, exist_ok=True)
            self.doc.save(str(ruta))
            self.logger.info("FIN guardar")
        except Exception as exc:
            self.logger.error(f"Error guardando documento en {ruta_archivo}: {exc}")
            raise
