import os
import pandas as pd
from datetime import datetime
import time
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

from conexion_config.conexion import ConexionAS400
from as400_core.ejecutor_cl import EjecutorCL
from jira.jira_utilidades import JiraClient

import warnings
from conexion_config.logging_config import get_logger

# Logger por módulo: mostrará el nombre del archivo entre corchetes
default_logger = get_logger(__file__)

# Agrega esta línea para ignorar la advertencia de SQLAlchemy
warnings.filterwarnings('ignore', category=UserWarning, message='.*pandas only supports SQLAlchemy connectable.*')

# --- IMPORTACIÓN DE LA FUNCIÓN DE CORREO ---
try:
    from utilidades.notificaciones_correo import NotificadorCorreo
except ImportError:
    pass


def obtener_conexion(logger=None):
    log = logger or default_logger
    try:
        conexion_as400 = ConexionAS400()
        conexion = conexion_as400.conectar()
        if not conexion:
            raise ConnectionError("La conexión al AS400 retornó nula.")
        return conexion
    except Exception as e:
        log.error(f"Error crítico de conexión AS400: {e}")
        raise


# =========================================================
# EJECUCIÓN SEGURA DE QUERIES
# =========================================================
def ejecutar_query_seguro(conexion, query, params=None, logger=None):
    try:
        return pd.read_sql(query, conexion, params=params)
    except Exception as e:
        if logger:
            logger.error(f"Error ejecutando query: {query} | Params: {params} | Error: {e}")
        return pd.DataFrame()


# =========================================================
# OBTENER ÓRDENES
# =========================================================
def obtener_ordenes_krws(base, limite=None, logger=None):
    conexion = obtener_conexion(logger)
    try:
        query = f"""SELECT DISTINCT TRIM(ASN) AS ASN FROM {base}.KRWSPWASIN"""
        if limite:
            query += f" FETCH FIRST {limite} ROWS ONLY"
            #logger.info(f"query ejecutada:{query}")
        return ejecutar_query_seguro(conexion, query, logger=logger)
    finally:
        conexion.close()

def ejecutar_comando_cl(comando, submit_job=False, job_name=None):
    """
    Ejecuta un comando CL en AS400
    """
    try:
        cl = EjecutorCL()
        
        if submit_job:
            respuesta = cl.submit_job(comando, job_name=job_name)
        else:
            respuesta = cl.ejecutar(comando)
            
        return respuesta
    except Exception as e:
        print(f"Error al ejecutar comando CL: {e}")
        return None

# =========================================================
# INVENTARIOS
# =========================================================
def obtener_datos_inventario(conexion, tabla, base, numero_orden, logger=None):
    query = f"""
        SELECT INVENTARIO,STATUSRI,U48RE,SKU,POÑRE,STÑRE,DATERE,ASN,UDF3,LN4RE FROM {base}.{tabla} WHERE ASN = ?
    """
    return ejecutar_query_seguro(conexion, query, params=[numero_orden], logger=logger)

def obtener_ptypp(conexion, base,asn,logger):
    query = f"""
        SELECT U29PS,U48PS,p.SKUPS  FROM {base}.PTYPP p WHERE p.POÑPS IN (SELECT POÑRE FROM {base}.KRWSPWINRC k WHERE K.ASN =? AND SKUPS = SUBSTRING(K.SKURE,2,6))
    """
    return ejecutar_query_seguro(conexion, query, params=[asn], logger=logger)

# =========================================================
# VALIDACIÓN (MOVIMIENTOS)
# =========================================================


def validar_inventario_incremental(df_inv_inicial, df_inv_final, ptypp, numero_orden, logger=None):
    import pandas as pd

    # =====================================================
    # 1. ESTRUCTURA DE RETORNO
    # =====================================================
    resultado = {
        "orden": numero_orden,
        "estado": "OK",
        "errores": [],
        "registros_validados": 0,
        "df_ini_detalle": pd.DataFrame(),
        "df_fin_detalle": pd.DataFrame(),
        "df_analisis": pd.DataFrame()
    }

    if df_inv_inicial.empty and df_inv_final.empty:
        resultado["estado"] = "ERROR"
        resultado["errores"].append("DataFrames vacíos.")
        return resultado

    # =====================================================
    # 2. NORMALIZACIÓN
    # =====================================================
    cols_visual = [
        "INVENTARIO", "STATUSRI", "U48RE", "SKU",
        "POÑRE", "STÑRE", "DATERE", "ASN", "UDF3", "LN4RE"
    ]

    cols_clave_agrupacion = ["SKU", "ASN", "POÑRE"]
    cols_suma = ["INVENTARIO", "U48RE", "UDF3"]

    def normalizar_df(df_input):
        if df_input.empty:
            return pd.DataFrame(columns=cols_visual)

        df = df_input.copy()

        for col in cols_visual:
            if col not in df.columns:
                df[col] = "N/A"

        for col in ["SKU", "ASN"]:
            df[col] = (
                df[col].astype(str)
                .str.strip()
                .str.upper()
                .replace("NAN", "")
            )

        df["POÑRE"] = pd.to_numeric(df["POÑRE"], errors="coerce").fillna(0).astype(int)

        for col in cols_suma:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0).astype(int)

        df["LN4RE"] = pd.to_numeric(df["LN4RE"], errors="coerce").fillna(0).astype(int)

        return df[cols_visual]

    df_ini_clean = normalizar_df(df_inv_inicial)
    df_fin_clean = normalizar_df(df_inv_final)

    resultado["df_ini_detalle"] = df_ini_clean
    resultado["df_fin_detalle"] = df_fin_clean

    # =====================================================
    # 3. PREPARAR PTYPP
    # =====================================================
    if ptypp is not None and not ptypp.empty:
        if "SKUPS" in ptypp.columns:
            ptypp = ptypp.rename(columns={"SKUPS": "SKU"})

        ptypp["SKU"] = ptypp["SKU"].astype(str).str.strip().str.upper()
        ptypp["U29PS"] = pd.to_numeric(ptypp["U29PS"], errors="coerce").fillna(0).astype(int)
        ptypp["U48PS"] = pd.to_numeric(ptypp["U48PS"], errors="coerce").fillna(0).astype(int)

        ptypp_clean = (
            ptypp.groupby("SKU")[["U29PS", "U48PS"]]
            .sum()
            .reset_index()
        )
    else:
        ptypp_clean = pd.DataFrame(columns=["SKU", "U29PS", "U48PS"])

    # =====================================================
    # 4. AGRUPACIÓN INTELIGENTE (SIN LN)
    # =====================================================
    def agrupar_datos(df):
        df_sum = (
            df.groupby(cols_clave_agrupacion, dropna=False)[cols_suma]
            .sum()
            .reset_index()
        )

        df_ln = (
            df.groupby(cols_clave_agrupacion)["LN4RE"]
            .first()
            .reset_index()
        )

        return df_sum.merge(df_ln, on=cols_clave_agrupacion, how="left")

    ini_group = agrupar_datos(df_ini_clean).rename(columns={
        "INVENTARIO": "INVENTARIO_INI",
        "U48RE": "U48RE_INI",
        "UDF3": "UDF3_INI"
    })

    fin_group = agrupar_datos(df_fin_clean).rename(columns={
        "INVENTARIO": "INVENTARIO_FINAL"
    })

    df_analisis = ini_group.merge(
        fin_group[["SKU", "ASN", "POÑRE", "INVENTARIO_FINAL"]],
        on=cols_clave_agrupacion,
        how="outer"
    )

    df_analisis = df_analisis.merge(ptypp_clean, on="SKU", how="left")

    # =====================================================
    # 5. LIMPIEZA POST-MERGE
    # =====================================================
    for col in [
        "INVENTARIO_INI", "U48RE_INI", "UDF3_INI",
        "INVENTARIO_FINAL", "U29PS", "U48PS"
    ]:
        df_analisis[col] = df_analisis[col].fillna(0).astype(int)

    for col in ["SKU", "ASN"]:
        df_analisis[col] = df_analisis[col].fillna("?").astype(str)

    df_analisis["LN4RE"] = df_analisis["LN4RE"].fillna("?").astype(str)

    # =====================================================
    # 6. CÁLCULOS
    # =====================================================
    df_analisis["ESPERADO"] = (
        df_analisis["INVENTARIO_INI"] + df_analisis["U48RE_INI"]
    )

    df_analisis["ESTADO"] = "OK"
    df_analisis["OBSERVACION"] = ""

    total_lineas = len(df_analisis)

    es_orden_cancelada = (
        total_lineas > 0 and
        len(df_analisis[
            (df_analisis["U48RE_INI"] == 0) &
            (df_analisis["UDF3_INI"] == 0)
        ]) == total_lineas
    )

    es_orden_redireccionada = (
        total_lineas > 0 and
        len(df_analisis[df_analisis["UDF3_INI"] != 0]) == total_lineas
    )

    # =====================================================
    # 7. VALIDACIÓN FINAL (REGLAS + MATEMÁTICA)
    # =====================================================
    for idx, row in df_analisis.iterrows():
        observaciones = []

        # --- Reglas de negocio ---
        if es_orden_cancelada:
            observaciones.append("ORDEN CANCELADA")
        elif row["U48RE_INI"] == 0 and row["UDF3_INI"] == 0:
            observaciones.append("LÍNEA CANCELADA")

        if es_orden_redireccionada:
            observaciones.append("ORDEN REDIRECCIONADA")
        elif row["UDF3_INI"] != 0:
            observaciones.append(f"REDIRECCIONADA ({row['UDF3_INI']})")

        if row["U48PS"] > 0 and row["U29PS"] < row["U48PS"]:
            observaciones.append("RECEPCION PARCIAL")

        if row["LN4RE"] == "?":
            observaciones.append("LN NO DISPONIBLE")

        # --- Validación matemática ---
        if row["INVENTARIO_FINAL"] != row["ESPERADO"]:
            df_analisis.at[idx, "ESTADO"] = "ERROR"
            resultado["estado"] = "ERROR"

            resultado["errores"].append(
                f"SKU {row['SKU']}: "
                f"Ini({row['INVENTARIO_INI']}) + "
                f"Carga({row['U48RE_INI']}) = "
                f"Esp({row['ESPERADO']}) vs "
                f"Real({row['INVENTARIO_FINAL']})"
            )

        # --- Asignación final ---
        if observaciones:
            df_analisis.at[idx, "OBSERVACION"] = " | ".join(dict.fromkeys(observaciones))
        else:
            df_analisis.at[idx, "OBSERVACION"] = "SIN OBSERVACIONES"

        resultado["registros_validados"] += 1

    resultado["df_analisis"] = df_analisis
    return resultado

# REPORTE WORD
# =========================================================



def generar_reporte_validacion_word(resultado_validacion, output_dir="reportes"):
    os.makedirs(output_dir, exist_ok=True)
    orden = resultado_validacion["orden"]
    nombre_archivo = f"REPORTE_{orden}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    ruta = os.path.join(output_dir, nombre_archivo)
    
    doc = Document()
    
    # Márgenes
    section = doc.sections[0]
    section.left_margin = Inches(0.2)
    section.right_margin = Inches(0.2)

    # Título
    t = doc.add_heading(f"REPORTE VALIDACIÓN: {orden}", level=1)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Estado
    p = doc.add_paragraph()
    run = p.add_run(f"Estado Global: {resultado_validacion['estado']}")
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 0, 0) if resultado_validacion['estado'] == "ERROR" else RGBColor(0, 128, 0)

    def crear_tabla(df, titulo, columnas_mostrar=None):
        doc.add_heading(titulo, level=2)
        if df is None or df.empty:
            doc.add_paragraph("Sin datos.")
            return

        cols = columnas_mostrar if columnas_mostrar else df.columns
        cols = [c for c in cols if c in df.columns]

        t = doc.add_table(rows=1, cols=len(cols))
        t.style = 'Table Grid'
        t.autofit = False
        
        # Headers
        hdr = t.rows[0].cells
        for i, col in enumerate(cols):
            texto = str(col).replace("_INI", "").replace("_FINAL", "")
            hdr[i].text = texto
            run = hdr[i].paragraphs[0].runs[0]
            run.font.bold = True
            run.font.size = Pt(7)
        
        # Data
        for _, row in df.iterrows():
            row_cells = t.add_row().cells
            for i, col in enumerate(cols):
                row_cells[i].text = str(row[col])
                row_cells[i].paragraphs[0].runs[0].font.size = Pt(7)

    # 1. Detalle Inicial y Final
    crear_tabla(resultado_validacion["df_ini_detalle"], "1. Inventario Inicial")
    crear_tabla(resultado_validacion["df_fin_detalle"], "2. Inventario Final")

    # 3. ANÁLISIS DETALLADO
    # Ordenamos las columnas para facilitar la lectura de las causas
    cols_analisis = [
        "SKU", "ASN", 
        "INVENTARIO_INI", "U48RE_INI", "ESPERADO", "INVENTARIO_FINAL", # Matemática
        "U29PS", "U48PS", "UDF3_INI", # Datos de control para las causas
        "ESTADO", 
        "OBSERVACION" # El texto explicativo
    ]
    
    crear_tabla(resultado_validacion["df_analisis"], "3. Análisis:", columnas_mostrar=cols_analisis)

    # 4. RESUMEN EJECUTIVO
    doc.add_heading("4. Resumen Global", level=2)
    df_a = resultado_validacion["df_analisis"]
    if df_a is not None and not df_a.empty:
        obs_list = [o for o in df_a["OBSERVACION"].unique() if o and str(o).strip() != ""]
        txt_obs = "\n".join(obs_list) if obs_list else "Flujo Normal (Sin incidencias)"
        
        totals = [
            ("Total Inventario Inicial", df_a["INVENTARIO_INI"].sum()),
            ("Total Recibido (U48RE)", df_a["U48RE_INI"].sum()),
            ("Total Esperado", df_a["ESPERADO"].sum()),
            ("Total Final", df_a["INVENTARIO_FINAL"].sum()),
            ("Observación", txt_obs)
        ]
        
        t_res = doc.add_table(rows=len(totals), cols=2)
        t_res.style = 'Table Grid'
        for i, (k, v) in enumerate(totals):
            c = t_res.rows[i].cells
            c[0].text = str(k)
            c[0].paragraphs[0].runs[0].font.bold = True
            c[1].text = str(v)
            if i == 4: # Resaltar las causas
                c[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 255)

    doc.save(ruta)

    ruta_relativa = os.path.relpath(ruta)
    tipo ='PO'
    jira = JiraClient()
    jira.main_jira(orden, ruta_relativa, tipo, "RI12DB") # Asumiendo RI12DB por defecto anterior
    return ruta_relativa

def ejecutar_stored_procedure_PO(conexion, procedure_call):
    """Ejecuta un stored procedure"""
    try:
        cursor = conexion.cursor()
        cursor.execute(procedure_call)
        conexion.commit()
        return True
    except Exception as e:
        print(f"Error al ejecutar stored procedure: {e}")
        return False
def obtener_conteo_registros(base,conexion, tabla, numero_orden,logger):
    """
    Obtiene el conteo de registros en una tabla para una orden específica
    """
    try:
        query = f"SELECT COUNT(*) as CONTEO FROM {base}.{tabla} WHERE ASN = '{numero_orden}'"
        #logger.info(f"query ejecutada: {query}")
        cursor = conexion.cursor()
        cursor.execute(query)
        resultado = cursor.fetchone()
        return resultado[0] if resultado else 0
    except Exception as e:
        # print(f"Error al obtener conteo de {tabla}: {e}") # Comentado para evitar flood en consola
        return -1
    
def normalizar_conteo(valor):
    """
    Normaliza cualquier valor de conteo a entero.
    Retorna 0 cuando el valor es None, vacío o inválido.
    """
    if valor is None:
        return 0

    if isinstance(valor, str):
        valor = valor.strip()
        if valor == "":
            return 0

    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return 0

def esperar_finalizacion_trabajoPO(base,numero_orden,conexion,logger,tabla_origen,tabla_destino,job_name=None,timeout_minutos=10):
    """
    Espera a que un proceso batch finalice validando movimientos
    entre tablas dinámicas (multi–base / multi–tabla).
    """

    # ================================
    # NORMALIZACIÓN DE TIMEOUT
    # ================================
    try:
        timeout_minutos = int(timeout_minutos)
    except (TypeError, ValueError):
        logger.warning(
            f"timeout_minutos inválido ({timeout_minutos}), usando valor por defecto: 10"
        )
        timeout_minutos = 10

    inicio_monitoreo = time.time()
    timeout_segundos = timeout_minutos * 60
    intervalo_verificacion = 5

    logger.info(
        f"Iniciando monitoreo | Orden={numero_orden} | "
        f"Origen={tabla_origen} a Destino={tabla_destino} | "
        f"Timeout={timeout_minutos} min"
    )

    try:
        # ================================
        # CONTEO INICIAL
        # ================================
        logger.info("Obteniendo conteo inicial de tablas")

        conteo_inicial_origen = normalizar_conteo(
            obtener_conteo_registros(
                base, conexion, tabla_origen, numero_orden, logger
            )
        )

        conteo_inicial_destino = normalizar_conteo(
            obtener_conteo_registros(
                base, conexion, tabla_destino, numero_orden, logger
            )
        )

        logger.info(
            f"Estado inicial | "
            f"{tabla_origen}: {conteo_inicial_origen}, "
            f"{tabla_destino}: {conteo_inicial_destino}"
        )

        if conteo_inicial_origen < 0 or conteo_inicial_destino < 0:
            logger.warning("Conteo inicial inválido. Se asume proceso exitoso.")
            time.sleep(30)
            return True

        if conteo_inicial_origen == 0:
            logger.info("Tabla origen ya vacía. Proceso completado.")
            return True

        # ================================
        # LOOP DE MONITOREO
        # ================================
        ultimo_conteo_origen = conteo_inicial_origen
        tiempo_sin_cambios = 0

        while (time.time() - inicio_monitoreo) < timeout_segundos:
            time.sleep(intervalo_verificacion)

            conteo_actual_origen = normalizar_conteo(
                obtener_conteo_registros(
                    base, conexion, tabla_origen, numero_orden, logger
                )
            )

            conteo_actual_destino = normalizar_conteo(
                obtener_conteo_registros(
                    base, conexion, tabla_destino, numero_orden, logger
                )
            )

            logger.info(
                f"Verificación | "
                f"{tabla_origen}: {conteo_actual_origen}, "
                f"{tabla_destino}: {conteo_actual_destino}"
            )

            # ================================
            # FINALIZACIÓN CORRECTA
            # ================================
            if (
                conteo_actual_origen == 0
                and conteo_actual_destino > conteo_inicial_destino
            ):
                logger.info("[OK] Registros procesados completamente.")
                return True

            # ================================
            # PROGRESO DETECTADO
            # ================================
            if conteo_actual_origen < ultimo_conteo_origen:
                logger.info(
                    f"[PROGRESO] {ultimo_conteo_origen} a {conteo_actual_origen}"
                )
                ultimo_conteo_origen = conteo_actual_origen
                tiempo_sin_cambios = 0
            else:
                tiempo_sin_cambios += intervalo_verificacion

            # ================================
            # VALIDACIÓN DE JOB (OPCIONAL)
            # ================================
            if tiempo_sin_cambios >= 60 and job_name:
                logger.info(f"[VERIFY] Verificando job activo: {job_name}")
                try:
                    activo = verificar_trabajos_activos(conexion, job_name)
                    if activo is False:
                        logger.warning(
                            "No hay job activo y no hay progreso en los registros."
                        )
                        return False
                except Exception as e:
                    logger.debug(f"No se pudo validar job: {e}")

                tiempo_sin_cambios = 0

        # ================================
        # TIMEOUT
        # ================================
        logger.warning(
            f"Timeout alcanzado ({timeout_minutos} minutos) "
            f"para la orden {numero_orden}"
        )
        return False

    except Exception:
        logger.error("Error crítico en monitoreo", exc_info=True)
        time.sleep(30)
        return True



def verificar_trabajos_activos(conexion, nombre_trabajo_patron):
    """
    Verifica si hay trabajos activos que coincidan con el patrón especificado
    """
    try:
        # Consultar tabla de trabajos activos del sistema
        query = f"""
            SELECT COUNT(*) as CONTEO 
            FROM TABLE(QSYS2.ACTIVE_JOB_INFO()) 
            WHERE JOB_NAME LIKE '%{nombre_trabajo_patron}%'
            AND JOB_STATUS = 'ACTIVE'
        """
        cursor = conexion.cursor()
        cursor.execute(query)
        resultado = cursor.fetchone()
        return (resultado[0] > 0) if resultado else False
    except Exception as e:
        # print(f"Error al verificar trabajos activos: {e}") # Comentado para evitar flood en consola
        return None
# =========================================================
# MAIN POR ORDEN
# =========================================================

def enviar_correo_resumen_validacion(lista_resultados, logger=None):
    """
    Recibe una lista de diccionarios (retornados por procesar_orden_krws)
    y envía un único correo con la tabla resumen y todos los adjuntos.
    """
    if not lista_resultados:
        return False

    fecha_hoy = datetime.now().strftime('%Y-%m-%d')
    
    # Contadores
    total = len(lista_resultados)
    errores = sum(1 for r in lista_resultados if r['estado'] != 'OK')
    
    # Preparar adjuntos (solo rutas válidas)
    adjuntos_totales = [r['ruta_reporte'] for r in lista_resultados if r.get('ruta_reporte')]

    # Generar filas de la tabla HTML
    filas_html = ""
    for res in lista_resultados:
        orden = res['orden']
        estado = res['estado']
        
        # Formatear observaciones
        obs_raw = res.get('observacion', [])
        if hasattr(obs_raw, 'tolist'): obs_raw = obs_raw.tolist() # Convertir numpy array si es necesario
        
        # Limpiar textos vacíos
        obs_str = ", ".join([str(x) for x in obs_raw if x and str(x).strip()])
        if not obs_str: obs_str = "Sin incidencias"

        color = "#D32F2F" if estado != "OK" else "#2E7D32" # Rojo/Verde
        
        filas_html += f"""
        <tr>
            <td style="padding:5px; border-bottom:1px solid #ddd;">{orden}</td>
            <td style="padding:5px; border-bottom:1px solid #ddd; font-weight:bold; color:{color};">{estado}</td>
            <td style="padding:5px; border-bottom:1px solid #ddd; font-size:12px;">{obs_str}</td>
        </tr>
        """

    # Asunto y Cuerpo
    asunto_estado = "CON ERRORES" if errores > 0 else "EXITOSO"
    asunto = f"[RESUMEN PO-WMS] Validación WMS-INFOR - {fecha_hoy} ({asunto_estado})"

    cuerpo = f"""
    <html>
    <body style="font-family: Arial, sans-serif; font-size: 14px;">
        <h3 style="color:#2E74B5;">Resumen de Ejecución</h3>
        <p>Se han procesado <b>{total}</b> órdenes.</p>
        
        <table style="width:100%; border-collapse: collapse; border: 1px solid #ccc;">
            <tr style="background-color: #f2f2f2; text-align: left;">
                <th style="padding:8px; border-bottom:2px solid #ccc;">Orden</th>
                <th style="padding:8px; border-bottom:2px solid #ccc;">Estado</th>
                <th style="padding:8px; border-bottom:2px solid #ccc;">Observaciones</th>
            </tr>
            {filas_html}
        </table>
        
        <p style="margin-top:15px; color:#666; font-size:12px;">
            * Se adjuntan los reportes detallados de cada orden.
        </p>
    </body>
    </html>
    """

    # Enviar
    if logger: logger.info(f"Enviando correo global con {len(adjuntos_totales)} adjuntos.")
    notificador = NotificadorCorreo(flujo='wms_infor')
    return notificador.enviar(asunto, cuerpo, adjuntos=adjuntos_totales)
def procesar_orden_krws(base, numero_orden, logger=None):
    log = logger or default_logger
    conexion = obtener_conexion(log)

    try:
        log.info(f"--- Iniciando Procesamiento Orden: {numero_orden} ---")

        # 1. EJECUCIÓN DE PROCESOS (SP y Jobs)
        ejecutar_stored_procedure_PO(
            conexion, f"CALL ELOPEZ.PROCESAR_PO('{base}', '{numero_orden}')"
        )
        log.info(f"Procedimiento almacenado ejecutado: {base}-{numero_orden}")
        
        cod_pais = base[2:4]
        ejecutar_comando_cl(f"CALL PGM(RIUNICOM63/SIWINPO1CL) PARM('{cod_pais}')", submit_job=True, job_name=f"WMSI_PO{cod_pais}")
        log.info(f"Programa sometido")
        
        esperar_finalizacion_trabajoPO(base, numero_orden, conexion, log,tabla_origen="KRWSPWASIN", tabla_destino="KRWSPWINRC",job_name=f"WMSI_PO{cod_pais}", timeout_minutos="15")
        log.info(f"Tiempo de espera finalizada. Ejecutando segundo procedimiento...")
        
        ejecutar_stored_procedure_PO(
            conexion, f"CALL ELOPEZ.PROCESAR_PO('{base}', '{numero_orden}')"
        )
        log.info("Segundo procedimiento ejecutado")

        # 2. OBTENCIÓN DE DATOS
        df_inv_ini = obtener_datos_inventario(conexion, "INV_INICIAL", base, numero_orden, log)
        df_inv_fin = obtener_datos_inventario(conexion, "INV_FINAL", base, numero_orden, log)
        
        # Obtener PTYPP para validación de parciales
        ptypp = obtener_ptypp(conexion, base, numero_orden, log)

        # 3. VALIDACIÓN
        res_validacion = validar_inventario_incremental(
            df_inv_ini, df_inv_fin, ptypp, numero_orden, log
        )

        # 4. GENERACIÓN DE REPORTE
        ruta_reporte = generar_reporte_validacion_word(res_validacion)
        
        # IMPORTANTE: Convertir a ruta absoluta para que se guarde en la lista de adjuntos
        # y Outlook la encuentre al final del proceso global.
        ruta_absoluta = os.path.abspath(ruta_reporte)
        log.info(f"Reporte generado en: {ruta_absoluta}")

        # 5. RETORNO DE DATOS (SIN ENVIAR CORREO)
        # Extraemos las observaciones para facilitarle el trabajo a la función de correo global
        observaciones = []
        if res_validacion.get("df_analisis") is not None and not res_validacion["df_analisis"].empty:
            observaciones = res_validacion["df_analisis"]["OBSERVACION"].unique()

        return {
            "orden": numero_orden,
            "estado": res_validacion["estado"], # "OK" o "ERROR"
            "observacion": observaciones,       # Lista de causas (Cancelación, Parcial, etc.)
            "ruta_reporte": ruta_absoluta,      # Ruta lista para adjuntar
            "resultado_completo": res_validacion # Por si acaso se requiere más detalle
        }

    except Exception as e:
        log.error(f"Error procesando {numero_orden}: {e}")
        # En caso de error, retornamos None o un dict indicando fallo para que el loop no se rompa
        return {
            "orden": numero_orden,
            "estado": "FALLO_PROCESO",
            "observacion": [str(e)],
            "ruta_reporte": None
        }

    finally:
        conexion.close()
        log.info(f"--- Fin Procesamiento Orden: {numero_orden} ---")



def main_PO(BASE="RI14DB"):
    LIMITE = None
    
    logger = default_logger

    logger.info("=== INICIO PROCESO KRWS AUTOMÁTICO ===")
    
    # 1. Obtener la lista de órdenes
    try: 
        df_ordenes = obtener_ordenes_krws(BASE, LIMITE, logger)
    except Exception as e:
        logger.error(f"Error obteniendo lista de órdenes: {e}")
        # Creamos un DF vacío para que no falle el if siguiente
        df_ordenes = pd.DataFrame()

    # 2. Lógica de Procesamiento
    if df_ordenes.empty:
        # CASO A: No hay nada que hacer
        logger.info("No hay órdenes pendientes para procesar.")
        
        
    else:
        # CASO B: Hay órdenes
        logger.info(f"Se encontraron {len(df_ordenes)} órdenes pendientes.")
        
        resultados_acumulados = [] # Lista acumuladora para el reporte final

        for orden in df_ordenes["ASN"]:
            try:
                # Llamamos a la función (que ya NO envía correo, solo retorna datos)
                datos_orden = procesar_orden_krws(BASE, orden, logger)
                
                if datos_orden:
                    resultados_acumulados.append(datos_orden)
            except Exception:
                # El try/except dentro del for asegura que si una falla, las demás sigan
                logger.exception(f"Error procesando orden {orden}")

        # --- AL FINALIZAR EL BUCLE ---
        # 3. Enviar el correo global con todo lo acumulado
        if resultados_acumulados:
            logger.info(f"Enviando correo global con {len(resultados_acumulados)} resultados.")
            enviar_correo_resumen_validacion(resultados_acumulados, logger)
        else:
            logger.warning("Se intentaron procesar órdenes pero ninguna generó resultados válidos.")

    logger.info("=== FIN PROCESO KRWS AUTOMÁTICO ===")

if __name__ == "__main__":
    main_PO()