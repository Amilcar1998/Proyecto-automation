import subprocess
import sys
import os
import glob
import shutil
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font

# Función para instalar paquetes si no están presentes
def install_package(package):
    try:
        __import__(package)
    except ImportError:
        print(f"El paquete '{package}' no está instalado. Procediendo a instalarlo...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])
        print(f"El paquete '{package}' se ha instalado correctamente.")

# Verificar e instalar paquetes necesarios
required_packages = ['pandas', 'xlwt', 'openpyxl']
for package in required_packages:
    install_package(package)

# Importar paquetes después de asegurar que estén instalados
import pandas as pd
import xlwt

# Definir las carpetas a borrar
carpetas_a_borrar = ["local", "Importado"]

# Borrar las carpetas especificadas si existen
for carpeta in carpetas_a_borrar:
    if os.path.exists(carpeta):
        shutil.rmtree(carpeta)
        print(f"La carpeta '{carpeta}' ha sido borrada.")

# Crear una carpeta de respaldo con la fecha de hoy
fecha_hoy = datetime.now().strftime('%Y-%m-%d')
backup_dir = f'backup/{fecha_hoy}'

# Contar cuántas veces se ha ejecutado el programa hoy
ejecucion_num = 1
while os.path.exists(f'{backup_dir}/ejecucion_{ejecucion_num}'):
    ejecucion_num += 1

# Crear la carpeta de ejecución actual
ejecucion_dir = f'{backup_dir}/ejecucion_{ejecucion_num}'
os.makedirs(ejecucion_dir, exist_ok=True)

# Cargar el archivo de entrada (Excel)
df = pd.read_excel("Formatos/OC compras.xlsx")  # Cambia el nombre del archivo a tu archivo real

# Imprimir las columnas del DataFrame para verificar
print("Columnas disponibles en el DataFrame:")
print(df.columns.tolist())

# Borrar archivos existentes que inicien con "Plantilla"
for file in glob.glob("Plantillas/Plantilla*.xls"):  # Cambia a .xls
    os.remove(file)

# Crear un nuevo libro de trabajo para el resumen
resumen_wb = Workbook()
resumen_ws = resumen_wb.active
resumen_ws.title = "Resumen"

# Escribir encabezados en el resumen
resumen_ws['A1'] = 'ID'
resumen_ws['B1'] = 'TIPO'
resumen_ws['C1'] = 'PAIS'
resumen_ws['D1'] = 'PREDISTRIBUIDO'
resumen_ws['E1'] = 'Archivo Generado'
resumen_ws['E1'].font = Font(bold=True)

# Contador de filas en el resumen
resumen_fila = 2

# Agrupar los datos por 'ID', 'TIPO', 'PAIS' y 'PREDISTRIBUIDO'
for (id_valor, tipo_valor, pais_valor, predistribuido_valor), group in df.groupby(
        ['ID', 'TIPO', 'PAIS', 'PREDISTRIBUIDO']):

    # Determinar el directorio de salida basado en PAIS y PREDISTRIBUIDO
    if predistribuido_valor == 'SI':
        directory = f'Plantillas/{pais_valor}/SI/'
    else:
        directory = f'Plantillas/{pais_valor}/NO/'

    # Crear el directorio si no existe
    os.makedirs(directory, exist_ok=True)

    # Crear un nuevo libro de trabajo
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet('Datos')

    # Escribir encabezados
    sheet.write(0, 0, 'PAIS')
    sheet.write(0, 1, 'COMPANIA')
    sheet.write(0, 2, 'TIENDA')
    sheet.write(0, 3, 'SKU')
    sheet.write(0, 4, 'COLOR')
    sheet.write(0, 5, 'TALLA')
    sheet.write(0, 6, 'MISELANEO')
    sheet.write(0, 7, 'UNIDADES')
    sheet.write(0, 8, 'PRECIO')

    # Contador de filas, comenzamos desde 1 para que los datos empiecen desde la fila 2
    fila = 1

    # Escribir los datos a partir de la fila 2
    for _, row in group.iterrows():
        sheet.write(fila, 0, 4)  # PAIS
        sheet.write(fila, 1, 2)  # COMPANIA
        sheet.write(fila, 2, row['TIENDA'])  # TIENDA
        sheet.write(fila, 3, row['SKU'])  # SKU
        sheet.write(fila, 4, "")  # COLOR
        sheet.write(fila, 5, "")  # TALLA
        sheet.write(fila, 6, "")  # MISELANEO
        sheet.write(fila, 7, row['CANTIDADES'])  # UNIDADES
        sheet.write(fila, 8, "")  # PRECIO

        # Aumentar la fila para el siguiente registro
        fila += 1

    # Nombre del archivo de salida basado en ID y Predistribuido (sin 'ID_')
    file_name = f'{directory}{id_valor}_Plantilla_Predistribuido_{predistribuido_valor}.xls'  # Cambiado para eliminar 'ID_'

    # Guardar el libro de trabajo
    workbook.save(file_name)

    # Crear la estructura de carpetas para el backup
    backup_sub_dir = os.path.join(ejecucion_dir, pais_valor, predistribuido_valor)
    os.makedirs(backup_sub_dir, exist_ok=True)

    # Generar un nuevo nombre de archivo si ya existe
    backup_file_name = os.path.join(backup_sub_dir, os.path.basename(file_name))
    contador = 1
    while os.path.exists(backup_file_name):
        backup_file_name = os.path.join(backup_sub_dir,
                                        f"{id_valor}_Plantilla_Predistribuido_{predistribuido_valor}_{contador}.xls")
        contador += 1

    # Copiar el archivo generado a la carpeta de respaldo
    shutil.copy(file_name, backup_file_name)

    # Confirmación de creación de archivo
    print(f"Archivo {file_name} creado y copiado a {backup_file_name}.")

    # Escribir en el resumen
    resumen_ws.cell(row=resumen_fila, column=1, value=id_valor)
    resumen_ws.cell(row=resumen_fila, column=2, value=tipo_valor)
    resumen_ws.cell(row=resumen_fila, column=3, value=pais_valor)
    resumen_ws.cell(row=resumen_fila, column=4, value=predistribuido_valor)
    resumen_ws.cell(row=resumen_fila, column=5, value=f'=HYPERLINK("{file_name}", "Ver Archivo")')

    # Incrementar el contador de filas en el resumen
    resumen_fila += 1

# Guardar el archivo de resumen
resumen_file_name = f'Resumen_{fecha_hoy}.xlsx'
resumen_wb.save(resumen_file_name)
print(f"Resumen guardado en {resumen_file_name}.")