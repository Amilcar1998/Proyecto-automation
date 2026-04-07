# ============================================================
# main.py
# Conciliación AS400 KTRHP vs KPUHP
# Regla de negocio:
# - Llave única: SKU
# - KTRHP: SUM(U71UA) por SKU
# - KPUHP: SUM(U45UH) por SKU
# - Reporte Word con 3 tablas
# ============================================================

from conexion_config.conexion import ConexionAS400
from datetime import datetime, timedelta
from conexion_config.logging_config import get_logger
from docx import Document

# ============================================================
# CONFIGURACIÓN DE LOGGER
# ============================================================

logger = get_logger(__file__)

# ============================================================
# UTILIDADES
# ============================================================

def obtener_fecha_proceso(dias_atras: int = 1) -> str:
    """
    Retorna la fecha del proceso en formato YYYYMMDD
    """
    # Para pruebas puedes fijarla:
    fecha = "20260122"
    #fecha = (datetime.now() - timedelta(days=dias_atras)).strftime("%Y%m%d")

    if not fecha.isdigit() or len(fecha) != 8:
        raise ValueError(f"Fecha inválida generada: {fecha}")

    return fecha


def ejecutar_consulta(query: str, params: tuple, nombre_proceso: str) -> list:
    """
    Ejecuta una consulta SQL en AS400.
    Retorna lista de tuplas (nunca None).
    """
    conexion = None
    try:
        logger.info(f"Iniciando consulta [{nombre_proceso}]")
        conexion = ConexionAS400().conectar()

        with conexion.cursor() as cursor:
            cursor.execute(query, params)
            resultados = cursor.fetchall()

        logger.info(
            f"Consulta [{nombre_proceso}] finalizada. "
            f"Registros obtenidos: {len(resultados)}"
        )
        return resultados

    except Exception as e:
        logger.exception(f"Error ejecutando [{nombre_proceso}]: {e}")
        return []

    finally:
        if conexion:
            try:
                conexion.close()
            except Exception:
                pass

# ============================================================
# EXTRACCIÓN DE DATOS
# ============================================================

def extraer_ktrhp(base: str, fecha: str) -> list:
    """
    Extrae totales KTRHP por SKU (U71UA)
    """
    query = f"""
        SELECT
            SKUUA,
            SUM(U71UA) AS TOTAL_U71UA
        FROM {base}.KTRHP
        WHERE D88UA = ?
          AND C86UA = 'R'
        GROUP BY SKUUA
    """
    return ejecutar_consulta(query, (fecha,), "KTRHP_POR_SKU")


def extraer_kpuhp(base: str, fecha: str) -> list:
    
    """
    Extrae totales KPUHP por SKU (U45UH)
    """
    query = f"""
        SELECT
            SKUUH,
            SUM(U45UH) AS TOTAL_U45UH
        FROM {base}.KPUHP
        WHERE D88UH = ? AND JBDUH ='TRANSPO12'
        GROUP BY SKUUH
    """
    return ejecutar_consulta(query, (fecha,), "KPUHP_POR_SKU")

# ============================================================
# NORMALIZACIÓN A DICCIONARIOS
# ============================================================

def dict_ktrhp_por_sku(datos: list) -> dict:
    return {sku: (total or 0) for sku, total in datos}


def dict_kpuhp_por_sku(datos: list) -> dict:
    return {sku: (total or 0) for sku, total in datos}

# ============================================================
# ANÁLISIS / CONCILIACIÓN
# ============================================================

def analizar_por_sku(ktrhp: dict, kpuhp: dict) -> list:
    """
    Compara totales por SKU entre KTRHP y KPUHP
    """
    resultado = []

    for sku, total_ktrhp in ktrhp.items():
        total_kpuhp = kpuhp.get(sku, 0)

        diferencia = total_ktrhp - total_kpuhp

        if sku not in kpuhp:
            estado = "SIN DATOS"
        elif diferencia == 0:
            estado = "OK"
        else:
            estado = "ERROR"

        resultado.append({
            "SKU": sku,
            "TOTAL_KTRHP": total_ktrhp,
            "TOTAL_KPUHP": total_kpuhp,
            "DIFERENCIA": diferencia,
            "ESTADO": estado
        })

    return resultado

# ============================================================
# REPORTE WORD (3 TABLAS)
# ============================================================

def generar_reporte_word(
    ktrhp_por_sku: dict,
    kpuhp_por_sku: dict,
    analisis: list,
    fecha: str,
    ruta_salida: str
):
    doc = Document()

    doc.add_heading("Reporte de Conciliación AS400 – KTRHP vs KPUHP", level=1)
    doc.add_paragraph(f"Fecha de proceso: {fecha}")

    # ==================================================
    # TABLA 1 – KTRHP
    # ==================================================
    doc.add_heading("1. Resultados KTRHP (U71UA por SKU)", level=2)

    tabla_ktrhp = doc.add_table(rows=1, cols=2, style="Table Grid")
    hdr = tabla_ktrhp.rows[0].cells
    hdr[0].text = "SKU"
    hdr[1].text = "TOTAL U71UA"

    for sku, total in ktrhp_por_sku.items():
        row = tabla_ktrhp.add_row().cells
        row[0].text = str(sku)
        row[1].text = str(total)

    # ==================================================
    # TABLA 2 – KPUHP
    # ==================================================
    doc.add_heading("2. Resultados KPUHP (U45UH por SKU)", level=2)

    tabla_kpuhp = doc.add_table(rows=1, cols=2, style="Table Grid")
    hdr = tabla_kpuhp.rows[0].cells
    hdr[0].text = "SKU"
    hdr[1].text = "TOTAL U45UH"

    for sku, total in kpuhp_por_sku.items():
        row = tabla_kpuhp.add_row().cells
        row[0].text = str(sku)
        row[1].text = str(total)

    # ==================================================
    # TABLA 3 – ANÁLISIS
    # ==================================================
    doc.add_heading("3. Análisis de Conciliación por SKU", level=2)

    tabla_analisis = doc.add_table(rows=1, cols=5, style="Table Grid")
    hdr = tabla_analisis.rows[0].cells
    hdr[0].text = "SKU"
    hdr[1].text = "TOTAL KTRHP"
    hdr[2].text = "TOTAL KPUHP"
    hdr[3].text = "DIFERENCIA"
    hdr[4].text = "ESTADO"

    for r in analisis:
        row = tabla_analisis.add_row().cells
        row[0].text = str(r["SKU"])
        row[1].text = str(r["TOTAL_KTRHP"])
        row[2].text = str(r["TOTAL_KPUHP"])
        row[3].text = str(r["DIFERENCIA"])
        row[4].text = r["ESTADO"]

    doc.save(ruta_salida)
    logger.info(f"Reporte Word generado: {ruta_salida}")

# ============================================================
# MAIN
# ============================================================

def main_vendor():
    logger.info("=== INICIO CONCILIACIÓN AS400 ===")

    BASE = "RI12DB"
    fecha = obtener_fecha_proceso()

    datos_ktrhp = extraer_ktrhp(BASE, fecha)
    datos_kpuhp = extraer_kpuhp(BASE, fecha)

    ktrhp_por_sku = dict_ktrhp_por_sku(datos_ktrhp)
    kpuhp_por_sku = dict_kpuhp_por_sku(datos_kpuhp)

    analisis = analizar_por_sku(ktrhp_por_sku, kpuhp_por_sku)

    ruta_reporte = f"Reporte_Conciliacion_SKU_{fecha}.docx"
    generar_reporte_word(
        ktrhp_por_sku,
        kpuhp_por_sku,
        analisis,
        fecha,
        ruta_reporte
    )

    logger.info("=== FIN CONCILIACIÓN AS400 ===")

# ============================================================
# EJECUCIÓN
# ============================================================

if __name__ == "__main__":
    main_vendor()
