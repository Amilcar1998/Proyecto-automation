import os
import sys
import pandas as pd
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from conexion_config.conexion import ConexionAS400
from conexion_config.logging_config import get_logger
from wms_infor.transferencias_wms_infor import validar_inventario

logger = get_logger(__name__)

def ejecutar_query(con, q):
    cursor = con.cursor()
    cursor.execute(q)
    cols = [column[0] for column in cursor.description]
    return pd.DataFrame.from_records(cursor.fetchall(), columns=cols)

def generar_word_reporte_ejecutivo(resultados_validacion, numero_orden, reporte_path):
    doc = Document()
    
    # Título y estilos básicos
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    title = doc.add_heading(f'Reporte Ejecutivo de Transferencia - Orden {numero_orden}', 0)
    
    p_info = doc.add_paragraph()
    p_info.add_run(f'Fecha de generación: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}\n').bold = True
    
    estado_gral = resultados_validacion.get('estado', 'ERROR')
    p_estado = doc.add_paragraph()
    p_estado.add_run('Estado General de la Orden: ').bold = True
    if estado_gral == 'OK':
        run_est = p_estado.add_run('✅ TRANSFERENCIA EXITOSA Y CUADRADA')
        run_est.font.color.rgb = RGBColor(0, 128, 0)
    else:
        run_est = p_estado.add_run('❌ ERROR EN TRANSFERENCIA O DESCUADRE')
        run_est.font.color.rgb = RGBColor(255, 0, 0)
        
    doc.add_paragraph(f"Mensaje del sistema: {resultados_validacion.get('mensaje', '')}")
    doc.add_paragraph("-" * 80)

    doc.add_heading('Detalle de Movimientos por SKU', level=1)
    
    movimientos = resultados_validacion.get('movimientos_inventario', [])
    if not movimientos:
        doc.add_paragraph("No se detectaron movimientos de inventario.")
    else:
        for mov in movimientos:
            sku = mov.get('sku', '')
            origen = mov.get('origen', '')
            destino = mov.get('destino', '')
            unid_movidas = mov.get('unidades_movidas', 0)
            
            inv_origen_ini = mov.get('inv_origen', 0)
            inv_origen_fin_esperado = inv_origen_ini - unid_movidas
            
            inv_destino_ini = mov.get('inv_destino', 0)
            inv_destino_fin_esperado = inv_destino_ini + unid_movidas
            
            # Extraer reales de las tablas de validación si existen
            inv_origen_fin_real = inv_origen_fin_esperado
            inv_destino_fin_real = inv_destino_fin_esperado
            
            if 'tablas_validacion' in resultados_validacion:
                if 'origen' in resultados_validacion['tablas_validacion']:
                    df_o = resultados_validacion['tablas_validacion']['origen']
                    match_o = df_o[(df_o['SKU'] == sku) & (df_o['ORIGEN'] == origen)]
                    if not match_o.empty:
                        inv_origen_fin_real = int(match_o.iloc[0]['INV_FIN'])
                
                if 'destino' in resultados_validacion['tablas_validacion']:
                    df_d = resultados_validacion['tablas_validacion']['destino']
                    match_d = df_d[(df_d['SKU'] == sku) & (df_d['DESTINO'] == destino)]
                    if not match_d.empty:
                        inv_destino_fin_real = int(match_d.iloc[0]['INV_FIN'])
                        
            origen_ok = (inv_origen_fin_real == inv_origen_fin_esperado)
            destino_ok = (inv_destino_fin_real == inv_destino_fin_esperado)
            
            doc.add_heading(f'📦 Resumen de Movimiento: SKU {sku}', level=2)
            p_res = doc.add_paragraph()
            p_res.add_run('Se ordenó mover: ').bold = True
            p_res.add_run(f'{unid_movidas} unidades\n')
            p_res.add_run('De: ').bold = True
            p_res.add_run(f'Ubicación {origen} ')
            p_res.add_run('➔ Hacia: ').bold = True
            p_res.add_run(f'Ubicación {destino}')
            
            # ORIGEN
            doc.add_heading(f'1. ¿Qué pasó en el Origen ({origen})?', level=3)
            doc.add_paragraph(f'Inventario antes del movimiento: {inv_origen_ini}', style='List Bullet')
            doc.add_paragraph(f'Se retiraron: -{unid_movidas}', style='List Bullet')
            doc.add_paragraph(f'Inventario que debería quedar: {inv_origen_fin_esperado}', style='List Bullet')
            
            p_ori4 = doc.add_paragraph(style='List Bullet')
            p_ori4.add_run('Inventario real registrado: ').bold = True
            p_ori4.add_run(f'{inv_origen_fin_real} ')
            if origen_ok:
                run_ok = p_ori4.add_run('✅ (Cuadre perfecto)')
                run_ok.font.color.rgb = RGBColor(0, 128, 0)
            else:
                run_err = p_ori4.add_run('❌ (Descuadre)')
                run_err.font.color.rgb = RGBColor(255, 0, 0)

            # DESTINO
            doc.add_heading(f'2. ¿Qué pasó en el Destino ({destino})?', level=3)
            doc.add_paragraph(f'Inventario antes del movimiento: {inv_destino_ini}', style='List Bullet')
            doc.add_paragraph(f'Se recibieron: +{unid_movidas}', style='List Bullet')
            doc.add_paragraph(f'Inventario que debería quedar: {inv_destino_fin_esperado}', style='List Bullet')
            
            p_des4 = doc.add_paragraph(style='List Bullet')
            p_des4.add_run('Inventario real registrado: ').bold = True
            p_des4.add_run(f'{inv_destino_fin_real} ')
            if destino_ok:
                run_ok2 = p_des4.add_run('✅ (Cuadre perfecto)')
                run_ok2.font.color.rgb = RGBColor(0, 128, 0)
            else:
                run_err2 = p_des4.add_run('❌ (Descuadre)')
                run_err2.font.color.rgb = RGBColor(255, 0, 0)
                
            doc.add_paragraph("-" * 80)
            
    doc.save(reporte_path)

def generar_solo_reporte(numero_orden):
    BASE = 'RI14DB'
    logger.info(f"Generando reporte para {numero_orden} en {BASE}")
    
    conexion_obj = ConexionAS400()
    conexion = conexion_obj.conectar()
    if not conexion:
        logger.error("No se pudo conectar a AS400")
        return
        
    try:
        numero_orden_climpio = numero_orden.strip()

        query_orden = f"SELECT SFTTD, DOCTD, STTTD, SKUTD, U71TD, O43TD, D37TD, TIPO, DOCTDNEW, ASN, NUMORDEN FROM {BASE}.KTCHPWASIN WHERE TRIM(NUMORDEN) = '{numero_orden_climpio}'"
        df_orden = ejecutar_query(conexion, query_orden)

        query_inv_inicial = f"SELECT * FROM {BASE}.INVENTARIO_INICIAL WHERE TRIM(NUMORDEN) = '{numero_orden_climpio}' order by sku desc "
        df_inv_inicial = ejecutar_query(conexion, query_inv_inicial)
        
        query_inv_final = f"SELECT * FROM {BASE}.INVENTARIO_FINAL WHERE TRIM(NUMORDEN) = '{numero_orden_climpio}' order by sku desc "
        df_inv_final = ejecutar_query(conexion, query_inv_final)
        
        resultados_validacion = validar_inventario(df_inv_inicial, df_inv_final, df_orden, logger)
        
        reporte_dir = os.path.dirname(os.path.abspath(__file__))
        fecha_actual = datetime.now().strftime("%Y%m%d")
        reporte_nombre = f"Reporte_Orden_{numero_orden_climpio}_{fecha_actual}_NUEVO.docx"
        reporte_path = os.path.join(reporte_dir, reporte_nombre)
        
        generar_word_reporte_ejecutivo(resultados_validacion, numero_orden_climpio, reporte_path)
        
        print(f"\n[OK] Reporte NUEVO guardado en: {reporte_path}")
        
    finally:
        conexion.close()

if __name__ == "__main__":
    generar_solo_reporte("0000016436")
